import pathlib
import shutil
import openpyxl as xl


class FileWrongExtensionError(Exception): ...  # {candidate.name} is not an .xlsx file


class HeaderAppearsTwiceError(
    Exception
): ...  # The header {header} appears more than once


class NoHeaderValueError(Exception): ...  # No value found below the {header} header


class NoWorkbookSelectedError(Exception): ...  # No workbook has been selected


# FileNotFoundError # {candidate} does not exist


class MissingHeadersError(Exception):  # Missing header(s) in the workbook:
    def __init__(self, missing: list[str], *args):
        super().__init__(*args)

        self.missing: list[str] = missing
        # raise ValueError(
        #     "Missing header(s) in the workbook: "
        #     + ", ".join(f"`{name}`" for name in missing)
        # )


class WrongValueAmountsError(
    Exception
):  # The amount of file names and destinations aren't equal
    def __init__(
        self, file_name_auto: int, file_name_manual: int, destination_folder: int, *args
    ):
        super().__init__(*args)

        self.file_name_auto: int = file_name_auto
        self.file_name_manual: int = file_name_manual
        self.destination_folder: int = destination_folder

        # ValueError(
        #     "The amount of file names and destinations aren't equal "
        #     f"({len(self.file_name_auto)} / {len(self.file_name_manual)} / "
        #     f"{len(self.destination_folder)})."
        # )


class State:
    HEADERS = (
        "folder",
        "initial folder",
        "file name auto",
        "file name manual",
        "destination",
    )

    def __init__(self) -> None:
        self.workbook: pathlib.Path | None = None
        self.filetype: str = ".pdf"

        self.common_folder: str = ""
        self.initial_folder: str = ""

        self.file_name_auto: list[str] = []
        self.file_name_manual: list[str] = []
        self.destination_folder: list[str] = []

    @property
    def extension(self) -> str:
        filetype = (self.filetype or "").strip()

        if not filetype:
            return ".pdf"

        return filetype if filetype.startswith(".") else f".{filetype}"

    def set_workbook(self, path: str | pathlib.Path) -> None:
        candidate = pathlib.Path(path)

        if candidate.suffix.lower() != ".xlsx":
            raise FileWrongExtensionError(str(candidate.name))
        if not candidate.is_file():
            raise FileNotFoundError(str(candidate))

        self.workbook = candidate

    def parse_workbook(self) -> None:
        if self.workbook is None:
            raise NoWorkbookSelectedError()

        wb = xl.load_workbook(self.workbook, data_only=True)
        ws = wb.active

        found: dict[str, bool] = {name: False for name in self.HEADERS}

        self.common_folder = ""
        self.initial_folder = ""
        self.file_name_auto = []
        self.file_name_manual = []
        self.destination_folder = []

        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue

                header = str(cell.value).strip().replace("_", " ").strip().lower()
                if header not in found:
                    continue

                if found[header]:
                    raise HeaderAppearsTwiceError(str(header))

                found[header] = True

                if header in ("folder", "initial folder"):
                    value = ws.cell(row=cell.row + 1, column=cell.column).value
                    text = None if value is None else str(value).strip()

                    if not text:
                        raise NoHeaderValueError(str(header))

                    if header == "folder":
                        self.common_folder = text
                    else:
                        self.initial_folder = text

                    continue

                values: list[str] = []
                current = cell.row + 1

                while True:
                    value = ws.cell(row=current, column=cell.column).value
                    if value is None:
                        break

                    text = str(value).strip()
                    if not text:
                        break

                    values.append(text)
                    current += 1

                if header == "file name auto":
                    self.file_name_auto = values
                elif header == "file name manual":
                    self.file_name_manual = values
                else:
                    self.destination_folder = values

        wb.close()

        missing = [name for name, present in found.items() if not present]

        if missing:
            raise MissingHeadersError(missing)

        lengths = {
            len(self.file_name_auto),
            len(self.file_name_manual),
            len(self.destination_folder),
        }

        if len(lengths) != 1:
            raise WrongValueAmountsError(
                len(self.file_name_auto),
                len(self.file_name_manual),
                len(self.destination_folder),
            )

    def get_operations(self) -> list[tuple[pathlib.Path, pathlib.Path]]:
        extension = self.extension
        common = pathlib.Path(self.common_folder)

        return [
            (
                common / self.initial_folder / f"{auto}{extension}",
                common / destination / f"{manual}{extension}",
            )
            for auto, manual, destination in zip(
                self.file_name_auto,
                self.file_name_manual,
                self.destination_folder,
            )
        ]

    def copy_files(self) -> tuple[int, list[str]]:
        copied = 0
        problems: list[str] = []

        for source, destination in self.get_operations():
            try:
                destination.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(source, destination)
                copied += 1
            except Exception as exc:
                problems.append(f"{source.name} → {destination.name}: {exc}")

        return copied, problems
